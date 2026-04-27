from __future__ import annotations

from pathlib import Path

import pandas as pd
from openpyxl.styles import Font, PatternFill
from openpyxl.utils import get_column_letter

from cacao_aroma_pipeline.utils import ensure_dir, safe_sheet_name


HEADER_FILL = PatternFill(fill_type="solid", fgColor="1F4E78")
HEADER_FONT = Font(color="FFFFFF", bold=True)


def _infer_column_widths(df: pd.DataFrame, *, min_width: int, max_width: int) -> list[int]:
    widths: list[int] = []
    preview = df.head(200).copy()
    for idx, column in enumerate(df.columns):
        max_len = len(str(column))
        series = preview.iloc[:, idx].astype(str)
        if not series.empty:
            max_len = max(max_len, int(series.map(len).max()))
        widths.append(max(min_width, min(max_width, max_len + 2)))
    return widths


def write_workbook(
    workbook_path: Path,
    sheets: dict[str, pd.DataFrame],
    *,
    freeze_header: bool = True,
    autofilter: bool = True,
    min_width: int = 10,
    max_width: int = 60,
) -> Path:
    ensure_dir(workbook_path.parent)
    with pd.ExcelWriter(workbook_path, engine="openpyxl") as writer:
        for sheet_name, df in sheets.items():
            safe_name = safe_sheet_name(sheet_name)
            frame = df.copy()
            frame.to_excel(writer, sheet_name=safe_name, index=False)
            ws = writer.sheets[safe_name]
            if ws.max_row >= 1 and ws.max_column >= 1:
                for cell in ws[1]:
                    cell.fill = HEADER_FILL
                    cell.font = HEADER_FONT
            if freeze_header:
                ws.freeze_panes = "A2"
            if autofilter and ws.max_row >= 1 and ws.max_column >= 1:
                ws.auto_filter.ref = ws.dimensions
            for idx, width in enumerate(
                _infer_column_widths(frame, min_width=min_width, max_width=max_width),
                start=1,
            ):
                ws.column_dimensions[get_column_letter(idx)].width = width
    return workbook_path
