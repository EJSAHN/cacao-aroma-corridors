from __future__ import annotations
import json
from pathlib import Path
import tempfile
import unittest
import numpy as np
import pandas as pd
from openpyxl import Workbook, load_workbook
from cacao_marker_harmonization.sources.sources import load_associations, load_candidates, locate, load_resource
from cacao_marker_harmonization.io.tableio import TableBook
from cacao_marker_harmonization.sources.common import DataError
BASE = Path(__file__).resolve().parents[1]

class SourceTests(unittest.TestCase):

    def setUp(self):
        self.temp = tempfile.TemporaryDirectory()
        self.root = Path(self.temp.name)
        self.cfg = json.loads((BASE / 'config/study.json').read_text(encoding='utf-8'))['sources']

    def tearDown(self):
        self.temp.cleanup()

    def associations(self):
        wb = Workbook()
        s = wb.active
        s.title = 'synth_tot'
        h = ['Chromosome',
            'Position of the association peak (bp)', 'Position of haplotypic bloc start',
            'Position of haplotypic bloc end', 'N° haplotypic bloc', 'GWAS method ',
            'Sorting of marker ', 'Traits', 'p-value of the strongest association',
            'Explanation rate of the trait of the strongest association']
        s.append(['Title'])
        s.append(h)
        s.append([1, 100, 90, 110, 7, 'GLM', 'G7', 'aroma_R', 0.001, 0.2])
        s.append([1, 100, 90, 110, 7, 'GLM', 'MAF5', 'aroma_R', 0.002, 0.3])
        s = wb.create_sheet('synth_composés_fruité')
        s.append(h)
        s.append([1, 100, 90, 110, 7, 'GLM', 'G7', 'aroma_R', 0.001, 0.2])
        s = wb.create_sheet('synth_composés_non_fruité')
        s.append(h)
        s.append([1, 100, 7, 90, 110, 'GLM', 'G7', 'aroma_R', 0.001, 0.2])
        s = wb.create_sheet('other')
        s.append(h)
        s.append([1, 200, 190, 210, 8, 'GLM', 'G7', 'other_UR', 0.003, 0.4])
        p = self.root / 'a.xlsx'
        wb.save(p)
        return load_associations(p, self.cfg)

    def test_association_no_sheet_inflation(self):
        a = self.associations()
        self.assertEqual(len(a['raw_records']), 5)
        self.assertEqual(len(a['master_events']), 2)
        self.assertEqual(len(a['unique_peaks']), 1)

    def test_different_models_filters_preserved(self):
        a = self.associations()
        self.assertEqual(set(a['master_events'].marker_filter_reported), {'G7', 'MAF5'})

    def test_interval_rotation_only_reported(self):
        a = self.associations()
        r = a['interval_audit']
        self.assertIn('rotation_matches_master', set(r.master_relation))
        bad = a['raw_records'][a['raw_records'].source_sheet.eq('synth_composés_non_fruité')].iloc[0]
        self.assertEqual(bad.source_start_bp, 7)

    def test_extra_not_added_to_primary(self):
        a = self.associations()
        self.assertEqual(len(a['extra_records']), 1)
        self.assertEqual(len(a['union_peaks_sensitivity']), 2)

    def test_candidate_duplicate_union(self):
        wb = Workbook()
        ws = wb.active
        ws.title = 'Synth'
        ws.append(['Title'])
        ws.append(['Chromosome',
            'start', 'end', 'gene_id', 'gene_function', 'Monoterpene pathway',
            'Sugar pathway'])
        ws.append([1, 10, 20, 'g1', 'func', 'x', None])
        ws.append([1, 10, 20, 'g1', 'func2', None, 'x'])
        ws.append([1, 50, 60, 'g2', 'f', None, None])
        p = self.root / 'c.xlsx'
        wb.save(p)
        a = load_candidates(p, self.cfg)
        self.assertEqual(len(a['unique_genes']), 2)
        self.assertEqual(len(a['memberships']), 2)
        self.assertEqual(len(a['duplicates']), 2)

    def test_candidate_conflict_quarantined(self):
        wb = Workbook()
        ws = wb.active
        ws.title = 'Synth'
        ws.append(['gene_id', 'Chromosome', 'start', 'end'])
        ws.append(['g1', 1, 1, 10])
        ws.append(['g1', 2, 1, 10])
        p = self.root / 'c.xlsx'
        wb.save(p)
        a = load_candidates(p, self.cfg)
        self.assertEqual(len(a['coordinate_conflicts']), 2)
        self.assertEqual(len(a['unique_genes']), 0)

    def test_genotype_headers_not_samples(self):
        wb = Workbook()
        ws = wb.active
        ws.append(['rs#',
            'alleles', 'chrom', 'pos', 'strand', 'assembly#', 'center',
            'protLSID', 'assayLSID', 'panelLSID', 'QCcode', 's1', 's2'])
        ws.append(['m1', 'A/C', 1, 10, '+', 'v2', 'x', 'x', 'x', 'x', 'x', 'A', 'M'])
        p = self.root / 'g.xlsx'
        wb.save(p)
        r = load_resource(p, self.cfg['resources'][0], self.cfg)
        self.assertEqual(r.samples, ['s1', 's2'])
        self.assertEqual(r.genotypes.shape, (1, 2))

    def test_missing_input_fails(self):
        with self.assertRaises(DataError):
            locate(self.root, 'none.xlsx', ['none.xlsx'])

    def test_ambiguous_copies_fail(self):
        (self.root / 'one').mkdir()
        (self.root / 'two').mkdir()
        (self.root / 'one' / 'x.xlsx').write_bytes(b'one')
        (self.root / 'two' / 'x.xlsx').write_bytes(b'two')
        with self.assertRaises(DataError):
            locate(self.root, 'none/x.xlsx', ['x.xlsx'])

    def test_ooxml_unicode_and_na(self):
        wb = Workbook()
        ws = wb.active
        ws.append(['key', 'value'])
        ws.append(['α', 'NA'])
        ws.append(['thing', 42])
        p = self.root / 'r.xlsx'
        wb.save(p)
        with TableBook(p) as b:
            rr = list(b.rows(b.sheet_names[0]))
        self.assertEqual(rr[1][1], ['α', 'NA'])
        self.assertEqual(rr[2][1][1], 42)


if __name__ == '__main__':
    unittest.main()

class SourceEdgeTests(unittest.TestCase):

    def setUp(self):
        self.temp = tempfile.TemporaryDirectory()
        self.root = Path(self.temp.name)
        self.cfg = json.loads((BASE / 'config/study.json').read_text(encoding='utf-8'))['sources']

    def tearDown(self):
        self.temp.cleanup()

    def test_qualitative_pathway_flags_remain_separate(self):
        wb = Workbook()
        ws = wb.active
        ws.title = 'Synth'
        ws.append(['gene_id', 'Chromosome', 'start', 'end', 'Fatty acid pathway', 'Sugar pathway'])
        ws.append(['g1', 1, 10, 20, 'Near', '?'])
        ws.append(['g2', 1, 30, 40, 'x', 'x'])
        path = self.root / 'c.xlsx'
        wb.save(path)
        out = load_candidates(path, self.cfg)
        self.assertEqual(len(out['unique_genes']), 2)
        self.assertEqual(len(out['qualifier_annotations']), 2)
        self.assertEqual(set(out['memberships'].gene_id), {'g2'})

    def test_metadata_repeated_sample_id_is_not_new_sample(self):
        wb = Workbook()
        ws = wb.active
        ws.title = 'marker'
        ws.append(['marker name'])
        ws.append(['m1'])
        ds = wb.create_sheet('dna_sample')
        ds.append(['DNA sample ID', 'label'])
        ds.append(['s1', 'a'])
        ds.append(['s1', 'a'])
        ds.append(['s2', 'b'])
        st = wb.create_sheet('study')
        st.append(['Study'])
        st.append(['Auxiliary'])
        path = self.root / 'aux.xlsx'
        wb.save(path)
        r = load_resource(path, self.cfg['resources'][4], self.cfg)
        self.assertEqual(r.samples, ['s1', 's2'])
        self.assertTrue(any((m['field'] == 'repeated_sample_metadata_id' for m in r.metadata)))

    def test_excel_error_recorded_not_repaired(self):
        wb = Workbook()
        ws = wb.active
        ws.title = 'marker'
        ws.append(['SNP marker name', 'chromosome', 'snp position', 'remark marker'])
        ws.append(['m1', 1, 50, '#VALUE!'])
        path = self.root / 'm.xlsx'
        wb.save(path)
        r = load_resource(path, self.cfg['resources'][1], self.cfg)
        self.assertEqual(r.markers.iloc[0].remark, '#VALUE!')
        self.assertTrue(any((m['field'] == 'source_excel_error' for m in r.metadata)))
        self.assertTrue(r.markers.iloc[0].eligible_coordinate)

    def test_marker_coordinate_error_excluded_not_imputed(self):
        wb = Workbook()
        ws = wb.active
        ws.title = 'marker'
        ws.append(['SNP marker name', 'chromosome', 'snp position'])
        ws.append(['m1', 1, '#VALUE!'])
        path = self.root / 'm.xlsx'
        wb.save(path)
        r = load_resource(path, self.cfg['resources'][1], self.cfg)
        self.assertFalse(r.markers.iloc[0].eligible_coordinate)

    def test_duplicate_genotype_sample_columns_rejected(self):
        wb = Workbook()
        ws = wb.active
        ws.append(['marker_id', 'chrom', 'pos', 's1', 's1'])
        ws.append(['m1', 1, 20, 'A', 'C'])
        path = self.root / 'm.xlsx'
        wb.save(path)
        with self.assertRaises(DataError):
            load_resource(path, self.cfg['resources'][0], self.cfg)

    def test_conflicting_marker_id_coordinates_not_covered(self):
        wb = Workbook()
        ws = wb.active
        ws.append(['marker_id', 'chrom', 'pos', 's1', 's2'])
        ws.append(['m1', 1, 20, 'A', 'C'])
        ws.append(['m1', 2, 20, 'C', 'A'])
        path = self.root / 'm.xlsx'
        wb.save(path)
        r = load_resource(path, self.cfg['resources'][0], self.cfg)
        self.assertEqual(int(r.markers.eligible_coordinate.sum()), 0)

    def test_missing_formula_cache_is_error(self):
        wb = Workbook()
        ws = wb.active
        ws.append(['id', 'value'])
        ws.append(['x', '=1+1'])
        path = self.root / 'f.xlsx'
        wb.save(path)
        with self.assertRaises(DataError), TableBook(path) as book:
            list(book.rows(book.sheet_names[0]))
