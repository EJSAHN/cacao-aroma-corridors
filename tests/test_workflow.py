"""Raw-input orchestration tests. Native placement fixture is software-only."""
from __future__ import annotations
import argparse, copy, contextlib, hashlib, io, json, random, tempfile, unittest
from pathlib import Path
from unittest.mock import patch
from openpyxl import Workbook
from cacao_marker_harmonization import workflow as w
from cacao_marker_harmonization.configuration import load_study
from cacao_marker_harmonization.errors import DataError
from cacao_marker_harmonization.validation import compare_workbooks, same_value, verify_hash_manifest
from cacao_marker_harmonization.io.export import write_workbook

ROOT=Path(__file__).resolve().parents[1]

class WorkflowTests(unittest.TestCase):
    def test_missing_source_fails_without_output(self):
        with tempfile.TemporaryDirectory() as d:
            cfg=load_study(ROOT/'config/study.json')
            with self.assertRaises(DataError):w.source_paths(Path(d)/'missing',cfg)
            self.assertEqual(list(Path(d).iterdir()),[])

    def test_different_assembly_rejected(self):
        with tempfile.TemporaryDirectory() as d:
            cfg=json.loads((ROOT/'config/study.json').read_text(encoding='utf-8'));cfg['coverage']['required_assembly_accession']='OTHER'
            p=Path(d)/'study.json';p.write_text(json.dumps(cfg), encoding='utf-8')
            with self.assertRaises(DataError):load_study(p)

    def test_scratch_inside_source_rejected(self):
        with tempfile.TemporaryDirectory() as d:
            root=Path(d);args=argparse.Namespace(source_root=root/'raw',output_root=root/'out',scratch=root/'raw'/'tmp',cache_root=root/'cache',tools_root=root/'tools',reference_cache=root/'refcache',reference_fasta=None,reference_gff=None)
            with self.assertRaises(DataError):w.check_locations(args,{})

    def test_nested_cache_rejected(self):
        with tempfile.TemporaryDirectory() as d:
            root=Path(d);args=argparse.Namespace(source_root=root/'raw',output_root=root/'out',scratch=root/'tmp',cache_root=root/'cache',tools_root=root/'cache'/'tools',reference_cache=root/'refcache',reference_fasta=None,reference_gff=None)
            with self.assertRaises(DataError):w.check_locations(args,{})

    def test_numeric_agreement_is_not_missing_to_zero(self):
        self.assertFalse(same_value(None,0,1e-12));self.assertFalse(same_value(False,0,1e-12))
        self.assertTrue(same_value(.5,.5+1e-13,1e-12));self.assertFalse(same_value(.5,.501,1e-12))

    def test_manifest_traversal_rejected(self):
        with tempfile.TemporaryDirectory() as d:
            root=Path(d);(root/'OUTPUT_SHA256.json').write_text(json.dumps({'../out':'bad'}), encoding='utf-8')
            with self.assertRaises(DataError):verify_hash_manifest(root)

    def test_workbook_comparison_detects_changed_count(self):
        with tempfile.TemporaryDirectory() as d:
            a=Path(d)/'a.xlsx';b=Path(d)/'b.xlsx'
            write_workbook(a,{'records':[dict(id='x',n=1,f=.5)]},[])
            write_workbook(b,{'records':[dict(id='x',n=2,f=.5)]},[])
            rows,detail=compare_workbooks(a,b)
            self.assertEqual(sum(r['mismatched_cells'] for r in rows),1)
            self.assertEqual(detail[0]['column'],'n')

    def test_workbook_comparison_requires_all_sheets(self):
        with tempfile.TemporaryDirectory() as d:
            a=Path(d)/'a.xlsx';b=Path(d)/'b.xlsx'
            write_workbook(a,{'records':[dict(id='x')]},[])
            write_workbook(b,{'other':[dict(id='x')]},[])
            with self.assertRaises(DataError):compare_workbooks(a,b)

    def test_raw_workflow_without_prior_outputs(self):
        # Known-position miniature FASTA/GFF and workbooks, not biological data.
        # Only native executable placement is replaced by a deterministic fixture.
        with tempfile.TemporaryDirectory() as d:
            root=Path(d);raw=root/'raw';raw.mkdir();refs=root/'reference';refs.mkdir()
            cfg=json.loads((ROOT/'config/study.json').read_text(encoding='utf-8'));n=24;total=30000
            cfg['sources']['chromosomes']=['1'];cfg['source_evidence']['chromosomes']=['1']
            cfg['reference'].update(assembly_name='FIXTURE_ASSEMBLY',assembly_accession='FIXTURE_ACCESSION',expected_total_bases=total,chromosomes=['1'],chromosome_aliases={'I':'1'},minimum_cache_free_gib=0)
            cfg['coverage']['required_assembly_accession']='FIXTURE_ACCESSION'
            cfg['coverage']['rarefaction_replicates']=2
            cfg['coverage']['background'].update(replicates=3,length_bins=1,local_gene_density_bins=1)
            cfg['dosage_ld']['required_assembly_accession']='FIXTURE_ACCESSION';cfg['dosage_ld']['chromosomes']=['1']
            cfg['dosage_ld']['ld']['context_pairs_per_class']=2
            cfg['alignment']['min_cache_free_gib']=0
            rng=random.Random(818);genome=''.join(rng.choice('ACGT') for _ in range(total))
            fa=refs/'fixture.fa';fa.write_text('>I chromosome:FIXTURE_ASSEMBLY:I:1:30000:1\n'+genome+'\n', encoding='utf-8')
            cfg['reference']['fasta'].update(filename=fa.name,sha256=None)
            gf=refs/'fixture.gff3';gf.write_text('##gff-version 3\n#!genome-build FIXTURE_ASSEMBLY\n'+''.join(f'I\tx\tgene\t{200+i*240}\t{249+i*240}\t.\t+\t.\tID=gene:g{i};Name=g{i};biotype=protein_coding\n' for i in range(100)), encoding='utf-8')
            cfg['reference']['gff'].update(filename=gf.name,sha256=None)
            # Source association workbook with a duplicate subset and explicit blocks.
            wb=Workbook();ws=wb.active;ws.title=cfg['sources']['association_master_sheet']
            head=['Chromosome','Position of the association peak (bp)','Position of haplotypic bloc start','Position of haplotypic bloc end','N° haplotypic bloc','GWAS method','Sorting of marker','Traits','p-value of the strongest association','Explanation rate of the trait of the strongest association']
            ws.append(head)
            for i in range(3):ws.append([1,234+i*240,200+i*240,249+i*240,i+1,'GLM','G7','aroma_R',.001,.2])
            ws=wb.create_sheet(cfg['sources']['association_fruity_subset_sheet']);ws.append(head);ws.append([1,234,200,249,1,'GLM','G7','aroma_R',.001,.2])
            p=raw/'assoc.xlsx';wb.save(p);cfg['sources']['association_file']=p.name
            wb=Workbook();ws=wb.active;ws.title=cfg['sources']['candidate_sheet'];ws.append(['Chromosome','start','end','gene_id','gene_function','Monoterpene pathway'])
            for i in range(3):ws.append([1,200+i*240,249+i*240,'g'+str(i),'fixture','x'])
            p=raw/'cand.xlsx';wb.save(p);cfg['sources']['candidate_file']=p.name
            markers=[]
            wb=Workbook();ws=wb.active;ws.title=cfg['source_evidence']['diversity_marker_sheet']
            ws.append(['SNP marker name','reference sequence name','chromosome','snp position','5flank_sequence','variation','3flank_sequence','version','strand','remark marker'])
            rs=wb.create_sheet(cfg['source_evidence']['diversity_reference_sheet']);rs.append(['reference sequence name','Sequence'])
            for i in range(n):
                start=201+i*240;tag=genome[start-1:start+68];alt=next(b for b in 'ACGT' if b!=tag[33]);key=f'{9000000+i}|F|0--33'
                markers.append((key,start,tag,alt));ws.append([key,key,1,start+33,tag[:33],tag[33]+'>'+alt,tag[34:],'V2','+','']);rs.append([key,tag])
            p=raw/'diversity.xlsx';wb.save(p)
            for spec in cfg['sources']['resources']:
                if spec['id']=='Diversity':spec['path']=p.name
                elif spec['id'] in ('Nacional_G7','Nacional_MAF5','Amazonia'):
                    wb=Workbook();ws=wb.active;ws.title='Sheet1';spec['path']=spec['id']+'.xlsx'
                    samples=['SNA604','L17H53']+[f's{i}' for i in range(58)]
                    ws.append(['rs#','alleles','chrom','pos','strand']+samples)
                    for i,(key,start,tag,alt) in enumerate(markers):ws.append([key,'A/C',1,start,'+']+['AMC'[(j+i)%3] for j in range(60)])
                    wb.save(raw/spec['path'])
            cp=root/'study.json';cp.write_text(json.dumps(cfg), encoding='utf-8');cfg=load_study(cp)
            args=argparse.Namespace(config=cp,source_root=raw,output_root=root/'out',scratch=root/'scratch',cache_root=root/'cache',tools_root=root/'tools',reference_cache=root/'refcache',reference_dir=None,reference_fasta=fa,reference_gff=gf,bowtie2_dir=None,offline=True,threads=1)
            before={p:w.sha256(p) for p in [*raw.glob('*.xlsx'),fa,gf]}
            def native_fixture(seqrows,exact,seqinv,faasset,config,args,out,save,log):
                decisions=[]
                for row in exact:
                    self.assertEqual(row['mapping_status'],'UNIQUE_EXACT_PLACEMENT')
                    decisions.append(dict(query_id=row['query_id'],source_row=row['source_row'],reference_key=row['reference_key'],source_marker_id=seqrows[row['source_row']-2]['marker_id'],placement_status='RETAINED_PRIMARY_CHROMOSOME',accepted_for_mapping=True,mapped_sequence_id=row['mapped_sequence_id'],mapped_chromosome='1',mapped_snp_bp=row['mapped_snp_bp'],mapped_strand=row['mapped_strand'],tag_start_bp=row['tag_start_bp'],tag_end_bp=row['tag_end_bp'],assembly_base=row['assembly_base'],other_source_base=row['other_source_base'],ref_cigar='69M',alt_cigar='69M'))
                save('01_realignment_evidence.xlsx',{'marker_decisions':decisions})
                return decisions,dict(tags=len(exact),sam_records=0,mapping_statuses={'FIXTURE':len(exact)})
            with patch.object(w,'native_alignment',side_effect=native_fixture), contextlib.redirect_stdout(io.StringIO()):
                out=w.run(args,cfg)
            obj=json.loads((out/'analysis_summary.json').read_text(encoding='utf-8'));self.assertEqual(obj['status'],'COMPLETED_WITH_DECLARED_LIMITATIONS')
            self.assertTrue(out.with_suffix('.zip').is_file());verify_hash_manifest(out)
            self.assertEqual(before,{p:w.sha256(p) for p in before})
            self.assertTrue((out/'alignment_coverage/04_paired_coverage_effects.xlsx').is_file())
            self.assertTrue((out/'dosage_ld/04_shared_collection_comparison.xlsx').is_file())
